"""BeamSummary: a list of beam sections read from a sections table and a forces table.

The numbers quoted in the docstrings are the ones the review of the
single-table format found (a support and a midspan merged under one label, a
compression face lost on export, a design that failed its own check), each
recalculated with ``Node.check()`` / ``Node.design()``.
"""

import copy
import io
import math
import warnings
from pathlib import Path
from typing import Any, Optional

import pandas as pd
import pytest
from docx.oxml.ns import qn
from docx.shared import Emu

from mento import (
    Concrete_ACI_318_19,
    Concrete_CIRSOC_201_25,
    Concrete_EN_1992_2004,
    Forces,
    MPa,
    Node,
    RectangularBeam,
    SteelBar,
    cm,
    kN,
    kNm,
    mm,
    set_language,
    split_single_table,
)
from mento.beam_summary import BeamSummary
from mento.i18n import translate
from mento.reports.summaries import SUMMARY_FONT_SIZE
from mento.results import FAIL_MARK, PASS_MARK, VERDICT_COLUMN, DocumentBuilder
from mento.shear_wall_summary import ShearWallSummary
from mento.slab_summary import OneWaySlabSummary
from mento.summary_base import GoverningDemand, SectionVerdict
from mento.summary_tables import SummaryInputError, SummaryInputWarning
from tests.reports.summary_data import (
    GEOMETRY_UNITS,
    beams,
    forces,
    geometry_only,
    one_continuous_section,
    slabs,
    support_and_midspan,
    wall_forces,
    walls,
)

pytestmark = pytest.mark.filterwarnings("ignore::UserWarning")


# ============================================================================
# FIXTURES
# ============================================================================


@pytest.fixture
def sample_concrete() -> Concrete_ACI_318_19:
    return Concrete_ACI_318_19(name="C25", f_c=25 * MPa)


@pytest.fixture
def sample_steel() -> SteelBar:
    return SteelBar(name="ADN 420", f_y=420 * MPa)


@pytest.fixture
def h25() -> Concrete_ACI_318_19:
    return Concrete_ACI_318_19(name="H25", f_c=25 * MPa)


@pytest.fixture
def sample_input_dataframe() -> pd.DataFrame:
    """The single table of mento 1.4.0 the summary tests were written against: four beams, one row each."""
    data = {
        "Label": ["", "V101", "V102", "V103", "V104"],
        "Comb.": ["", "ELU 1", "ELU 2", "ELU 3", "ELU 4"],
        "b": ["cm", 20, 20, 20, 20],
        "h": ["cm", 50, 50, 50, 50],
        "cc": ["mm", 25, 25, 25, 25],
        "Nx": ["kN", 0, 0, 0, 0],
        "Vz": ["kN", 20, -50, 100, 100],
        "My": ["kNm", 0, -35, 40, 45],
        "ns": ["", 0, 1.0, 1.0, 1.0],
        "dbs": ["mm", 0, 6, 6, 6],
        "sl": ["cm", 0, 20, 20, 20],
        "n1": ["", 2.0, 2, 2.0, 2.0],
        "db1": ["mm", 12, 12, 12, 12],
        "n2": ["", 1.0, 1, 1.0, 0.0],
        "db2": ["mm", 10, 16, 10, 0],
        "n3": ["", 2.0, 0.0, 2.0, 0.0],
        "db3": ["mm", 12, 0, 16, 0],
        "n4": ["", 0, 0.0, 0, 0.0],
        "db4": ["mm", 0, 0, 0, 0],
    }
    return pd.DataFrame(data)


@pytest.fixture
def sample_input_with_nan() -> pd.DataFrame:
    """A 1.4.0 table with blank cells, among them a diameter (db2 of V101) with no bars."""
    data = {
        "Label": ["", "V101", "V102"],
        "Comb.": ["", "ELU 1", "ELU 2"],
        "b": ["cm", 20, 20],
        "h": ["cm", 50, 50],
        "cc": ["mm", 25, 25],
        "Nx": ["kN", None, 0],
        "Vz": ["kN", 20, None],
        "My": ["kNm", 0, -35],
        "ns": [None, None, 1.0],
        "dbs": ["mm", 0, 6],
        "sl": ["cm", None, 20],
        "n1": ["", 2.0, 2],
        "db1": ["mm", 12, 12],
        "n2": ["", None, 1],
        "db2": ["mm", 10, 16],
        "n3": ["", 2.0, 0.0],
        "db3": ["mm", 12, 0],
        "n4": ["", 0, 0.0],
        "db4": ["mm", 0, 0],
    }
    return pd.DataFrame(data)


@pytest.fixture
def beam_summary(
    sample_concrete: Concrete_ACI_318_19, sample_steel: SteelBar, sample_input_dataframe: pd.DataFrame
) -> BeamSummary:
    """The four beams of the 1.4.0 table, converted: each row a section of its own."""
    return BeamSummary(sample_concrete, sample_steel, *split_single_table(sample_input_dataframe, "beam"))


def _row(summary: BeamSummary, label: str) -> pd.Series:
    """The check() row of one section."""
    table = summary.check()
    return table[table["Beam"] == label].iloc[0]


# ============================================================================
# TWO TABLES: WHAT A SECTION IS, AND WHAT IT CARRIES
# ============================================================================


def test_a_support_and_a_midspan_are_two_sections(h25: Any, sample_steel: SteelBar) -> None:
    """Case A: the support's 3Ø16 are cut before the midspan, so they are two sections.

    The midspan's 2Ø32 under +170 kN·m are not tension-controlled with the
    2Ø8 above them: DCRb,bot 1.263 and ❌. Under one label with the support
    (the single table of the pull request) it passed at 0.934, counting the
    support bars as compression steel.
    """
    summary = BeamSummary(h25, sample_steel, *support_and_midspan())
    support, midspan = summary.check().iloc[1], summary.check().iloc[2]

    assert (support["Comb.,top"], support["Mu,top"], support["DCRb,top"]) == ("apoyo", -60.0, 0.804)
    assert (support["Comb.,v"], support["Vu"], support["DCRv"]) == ("apoyo", 80.0, 0.535)
    assert support[VERDICT_COLUMN] == PASS_MARK
    assert (midspan["Comb.,bot"], midspan["Mu,bot"], midspan["DCRb,bot"]) == ("tramo", 170.0, 1.263)
    assert midspan[VERDICT_COLUMN] == FAIL_MARK
    codes = [(w.code, w.face) for w in summary.warnings[("", "V9t")]]
    assert codes == [("not_tension_controlled", "bottom"), ("stirrup_spacing_exceeds_compression_support", None)]
    assert midspan["Warnings"] == "not_tension_controlled (bottom), stirrup_spacing_exceeds_compression_support"


def test_a_label_given_twice_in_sections_raises(h25: Any, sample_steel: SteelBar) -> None:
    sections = beams([{"Label": "V9"}, {"Label": "V9 "}], units=GEOMETRY_UNITS)
    with pytest.raises(SummaryInputError) as raised:
        BeamSummary(h25, sample_steel, sections, forces([{"Label": "V9", "Comb.": "C", "My": 10}]))
    assert raised.value.code == "duplicate_section"
    assert "'V9'" in str(raised.value) and "rows 3, 4" in str(raised.value)
    assert "a support and a midspan" in str(raised.value)


def test_continuous_bars_are_one_section_declared_once(h25: Any, sample_steel: SteelBar) -> None:
    """Case B: one row with both faces and two rows of forces: 0.804 / 0.934 / 0.548, as a hand-built node."""
    summary = BeamSummary(h25, sample_steel, *one_continuous_section())
    row = summary.check().iloc[1]

    assert (row["Comb.,top"], row["Mu,top"], row["DCRb,top"]) == ("apoyo", -60.0, 0.804)
    assert (row["Comb.,bot"], row["Mu,bot"], row["DCRb,bot"]) == ("tramo", 170.0, 0.934)
    assert (row["Comb.,v"], row["Vu"], row["DCRv"]) == ("apoyo", 80.0, 0.548)
    assert (row["Warnings"], row[VERDICT_COLUMN]) == ("-", PASS_MARK)

    beam = RectangularBeam(label="V9", concrete=h25, steel_bar=sample_steel, width=20 * cm, height=40 * cm, c_c=25 * mm)
    beam.set_transverse_rebar(1, 10 * mm, 17 * cm)
    beam.set_longitudinal_rebar_top(3, 16 * mm)
    beam.set_longitudinal_rebar_bot(2, 32 * mm)
    node = Node(
        beam,
        [Forces(label="apoyo", V_z=80 * kN, M_y=-60 * kNm), Forces(label="tramo", V_z=10 * kN, M_y=170 * kNm)],
    )
    node.check()
    assert summary.nodes[0].section.flexure_checks == beam.flexure_checks
    assert summary.nodes[0].section.shear_checks == beam.shear_checks


def test_check_follows_node_check_order(sample_concrete: Any, sample_steel: SteelBar) -> None:
    """25x70, 4Ø25 ++ 2Ø20 below, 2Ø12 on top, 1eØ8/25, Mu 500, Vu 80: the bracing of the compression bars.

    Shear before flexure (the order of the pull request) gave DCRb,bot 0.922
    and ✅; the flexure check is what finds the face relied on as compression
    steel, and the shear check reads it.
    """
    sections = beams(
        [{"Label": "V", "n1_top": 2, "db1_top": 12, "n1_bot": 4, "db1_bot": 25, "n3_bot": 2, "db3_bot": 20}],
        b=25,
        h=70,
        legs=2,
        dbs=8,
        sl=25,
    )
    summary = BeamSummary(
        sample_concrete, sample_steel, sections, forces([{"Label": "V", "Comb.": "C1", "Vz": 80, "My": 500}])
    )
    row = summary.check().iloc[1]

    assert row["DCRb,bot"] == 0.922
    assert row[VERDICT_COLUMN] == FAIL_MARK
    assert [w.code for w in summary.results[0].warnings] == [
        "stirrup_spacing_exceeds_compression_support",
        "stirrup_diameter_below_compression_support",
    ]
    beam = RectangularBeam(
        label="V", concrete=sample_concrete, steel_bar=sample_steel, width=25 * cm, height=70 * cm, c_c=25 * mm
    )
    beam.set_transverse_rebar(1, 8 * mm, 25 * cm)
    beam.set_longitudinal_rebar_top(2, 12 * mm)
    beam.set_longitudinal_rebar_bot(4, 25 * mm, 0, None, 2, 20 * mm)
    node = Node(beam, [Forces(label="C1", V_z=80 * kN, M_y=500 * kNm)])
    node.check()
    assert summary.results[0].warnings == node.warnings


def test_capacity_check_leaves_the_verdict_alone(h25: Any, sample_steel: SteelBar) -> None:
    """The capacity check runs on copies: ``results`` and ``warnings`` stay those of the forces."""
    summary = BeamSummary(h25, sample_steel, *support_and_midspan())
    summary.check()
    results, found = summary.results, summary.warnings

    capacities = summary.check(capacity_check=True)
    summary.flexure_results(capacity_check=True)
    summary.shear_results(capacity_check=True)

    assert summary.results == results and summary.warnings == found
    assert ("not_tension_controlled", "bottom") in [(w.code, w.face) for w in summary.warnings[("", "V9t")]]
    assert list(capacities.columns[:6]) == ["Beam", "b×h", "As,top", "As,bot", "Av", "As,top,real"]
    assert capacities.iloc[1]["b×h"] == "20×40"
    assert {"ØMn,top", "ØMn,bot", "ØVn"} <= set(capacities.columns)


def test_a_one_sign_beam_reads_back_with_its_compression_bars(h25: Any, sample_steel: SteelBar, tmp_path: Path) -> None:
    """Case C: a design writes both faces, so export and import give back the section designed.

    20x40 under +170 designs 2Ø16 + 1Ø16 on top and 2Ø32 below, 0.934; with a
    row that could only write the face its moment pulls, the top came back as
    2Ø8 and the check as 1.263. A 20x50 under +260 kN·m: 0.975 before and
    after (it came back at 1.272).
    """
    summary = BeamSummary(h25, sample_steel, *geometry_only())
    designed = summary.design()
    support, midspan = designed.iloc[1], designed.iloc[2]
    assert (support["n1_top"], support["db1_top"], support["n2_top"], support["db2_top"]) == (2, 16, 1, 12)
    assert (support["n1_bot"], support["db1_bot"]) == (2, 10)
    assert (midspan["n1_top"], midspan["db1_top"], midspan["n2_top"], midspan["db2_top"]) == (2, 16, 1, 16)
    assert (midspan["n1_bot"], midspan["db1_bot"], midspan["legs"], midspan["dbs"], midspan["sl"]) == (2, 32, 2, 10, 17)
    before = summary.check()
    assert list(before["DCRb,top"][1:]) == [0.928, 0.0] and list(before["DCRb,bot"][1:]) == [0.0, 0.934]
    assert list(before[VERDICT_COLUMN][1:]) == [PASS_MARK, PASS_MARK]

    path = tmp_path / "beams.xlsx"
    summary.to_excel(path)
    again = BeamSummary.from_excel(h25, sample_steel, path)
    assert again.check().equals(before)
    assert again.results == summary.results

    deep = BeamSummary(
        h25,
        sample_steel,
        beams([{"Label": "V8"}], units=GEOMETRY_UNITS),
        forces([{"Label": "V8", "Comb.": "C1", "Vz": 60, "My": 260}]),
    )
    deep.design()
    deep.export_design(tmp_path / "deep.xlsx")
    deep.import_design(tmp_path / "deep.xlsx")
    assert deep.check().iloc[1]["DCRb,bot"] == 0.975


def test_a_section_with_no_forces_has_its_row(h25: Any, sample_steel: SteelBar, tmp_path: Path) -> None:
    sections = beams([{"Label": "V1", "n1_bot": 3, "db1_bot": 16}, {"Label": "V2", "n1_bot": 2, "db1_bot": 12}])
    with pytest.warns(SummaryInputWarning) as caught:
        summary = BeamSummary(h25, sample_steel, sections, forces([{"Label": "V1", "Comb.": "C", "Vz": 20, "My": 40}]))
    assert [w.message.code for w in caught] == ["section_without_forces"]
    assert "'V2'" in str(caught[0].message)
    assert [w.code for w in summary.input_warnings] == ["section_without_forces"]

    row = summary.check().iloc[2]
    assert row[VERDICT_COLUMN] == "no forces"
    assert math.isnan(row["DCRb,bot"]) and math.isnan(row["DCRv"])
    assert summary.results[1] == SectionVerdict(("", "V2"), "no_forces", None, None, None, (), None)
    # Kept as given, and written back.
    assert list(summary.sections_table["Label"][1:]) == ["V1", "V2"]
    summary.design()
    assert summary.sections_table.iloc[2]["n1_bot"] == 2
    with pytest.raises(SummaryInputError, match="no forces") as raised:
        summary.flexure_results(index="V2")
    assert raised.value.code == "no_forces_to_report"


def test_forces_of_an_unknown_section_raise(h25: Any, sample_steel: SteelBar) -> None:
    with pytest.raises(SummaryInputError) as raised:
        BeamSummary(
            h25,
            sample_steel,
            beams([{"Label": "V1"}], units=GEOMETRY_UNITS),
            forces([{"Label": "V1", "Comb.": "C", "My": 1}, {"Label": "V7", "Comb.": "C", "My": 1}]),
        )
    assert raised.value.code == "unknown_section"
    assert "'V7'" in str(raised.value) and "rows 4" in str(raised.value)


def _envelope_beams(order: int = 1) -> tuple:
    """V5, V6 and ±70 of the review: ACI 20x50, cc 25, 1eØ6/20."""
    sections = beams(
        [
            {"Label": "V5", "n1_bot": 2, "db1_bot": 12},
            {"Label": "V6", "n1_top": 3, "db1_top": 20, "n1_bot": 2, "db1_bot": 20},
            {"Label": "V70", "n1_top": 3, "db1_top": 16, "n1_bot": 3, "db1_bot": 16},
        ],
        legs=2,
        dbs=6,
        sl=20,
    )
    seventy = [{"Vz": 50, "My": 70}, {"Vz": 50, "My": -70}][::order]
    rows = [
        {"Label": "V5", "Comb.": "C1", "Vz": 100, "Nx": 300, "My": 20},
        {"Label": "V5", "Comb.": "C2", "Vz": 85, "Nx": -150, "My": 20},
        {"Label": "V6", "Comb.": "C1", "My": -160},
        {"Label": "V6", "Comb.": "C2", "My": 120},
        {"Label": "V70", "Comb.": "C1", **seventy[0]},
        {"Label": "V70", "Comb.": "C2", **seventy[1]},
    ]
    return sections, forces(rows)


def test_check_names_the_combination_of_each_dcr(sample_concrete: Any, sample_steel: SteelBar) -> None:
    """Each DCR next to the demand of the combination it came from.

    The envelope column by column put Vu 100 and Nu 300 of C1 next to the
    DCRv 1.025 of C2 (V5), and Mu -160 next to the DCRb,bot 1.181 of +120 (V6).
    """
    table = BeamSummary(sample_concrete, sample_steel, *_envelope_beams()).check()
    v5, v6, v70 = table.iloc[1], table.iloc[2], table.iloc[3]

    assert (v5["Comb.,v"], v5["Vu"], v5["Nu"], v5["DCRv"]) == ("C2", 85.0, -150.0, 1.025)
    assert (v6["Comb.,top"], v6["Mu,top"], v6["DCRb,top"]) == ("C1", -160.0, 1.089)
    assert (v6["Comb.,bot"], v6["Mu,bot"], v6["DCRb,bot"]) == ("C2", 120.0, 1.181)
    # Tied combinations are all named; the result does not depend on the order of the rows.
    assert (v70["Comb.,v"], v70["Vu"], v70["DCRv"]) == ("C1, C2", 50.0, 0.501)
    assert (v70["DCRb,top"], v70["DCRb,bot"]) == (0.712, 0.712)
    other = BeamSummary(sample_concrete, sample_steel, *_envelope_beams(order=-1)).check().iloc[3]
    assert other["Comb.,v"] == "C1, C2"
    assert (other["Mu,top"], other["Mu,bot"], other["DCRb,top"], other["DCRb,bot"]) == (-70.0, 70.0, 0.712, 0.712)


def test_the_named_combination_carries_that_dcr(sample_concrete: Any, sample_steel: SteelBar) -> None:
    """Over a corpus, the per-combination tables give the demand and the DCR ``check()`` names for each DCR."""
    corpus = [
        BeamSummary(sample_concrete, sample_steel, *_envelope_beams()),
        BeamSummary(Concrete_ACI_318_19(name="H25", f_c=25 * MPa), sample_steel, *one_continuous_section()),
        BeamSummary(Concrete_ACI_318_19(name="H25", f_c=25 * MPa), sample_steel, *support_and_midspan()),
    ]
    checked = 0
    for summary in corpus:
        summary.check()
        for record in summary.results:
            flexure = summary.flexure_results(index=record.label)
            shear = summary.shear_results(index=record.label)
            for demand, table, column, value_column in (
                (record.top, flexure, "Mu", "DCR"),
                (record.bottom, flexure, "Mu", "DCR"),
                (record.shear, shear, "Vu", "DCR"),
            ):
                assert demand is not None
                for name in demand.combinations:
                    rows = table[table["Comb."] == name.split(" (#")[0]]
                    assert any(
                        math.isclose(float(r[value_column]), round(demand.DCR, 3), abs_tol=1e-3)
                        or math.isclose(float(r[value_column]), demand.DCR, abs_tol=1e-3)
                        for _, r in rows.iterrows()
                    )
                    checked += 1
    assert checked > 10


def test_a_repeated_combination_raises(h25: Any, sample_steel: SteelBar) -> None:
    """A section takes each combination once: a name given twice for it stops the reading, naming both."""
    rows = [
        {"Label": "V9", "Comb.": "ENV", "Vz": 80, "My": 150},
        {"Label": "V9", "Comb.": "1.2D+1.6L", "Vz": 80, "My": -60},
        {"Label": "V9", "Comb.": "ENV", "Vz": 10, "My": 170},
    ]
    sections, _ = one_continuous_section()
    with pytest.raises(SummaryInputError) as raised:
        BeamSummary(h25, sample_steel, sections, forces(rows))
    assert raised.value.code == "duplicate_combination"
    assert "combination 'ENV' of 'V9'" in str(raised.value) and "rows 3, 5" in str(raised.value)
    # The same name on two sections is two combinations; no name at all is named by its position.
    other = beams([{"Label": "V9"}, {"Label": "V10"}], units=GEOMETRY_UNITS)
    rows = [{"Label": "V9", "Comb.": "ENV", "My": 5}, {"Label": "V10", "Comb.": "ENV", "My": 5}]
    rows += [{"Label": "V10", "Comb.": "", "My": 6}, {"Label": "V10", "Comb.": "", "My": 7}]
    summary = BeamSummary(h25, sample_steel, other, forces(rows))
    summary.design()
    assert summary.check().iloc[2]["Comb.,bot"] == "#3"


def test_check_dcr_columns_are_float(h25: Any, sample_steel: SteelBar) -> None:
    summary = BeamSummary(h25, sample_steel, *support_and_midspan())
    table = summary.check()
    data = table.iloc[1:]
    for column in ("DCRb,top", "DCRb,bot", "DCRv"):
        assert all(isinstance(value, float) for value in data[column])
    assert data.iloc[1]["DCRb,bot"] == 1.263
    assert "not_tension_controlled (bottom)" in data.iloc[1]["Warnings"]
    table.to_json()


def test_results_by_label(h25: Any, sample_steel: SteelBar) -> None:
    summary = BeamSummary(h25, sample_steel, *support_and_midspan())
    with pytest.raises(ValueError, match="either 'shear' or 'flexure'"):
        summary._process_beam_for_check(summary.nodes[0], "torsion", False)
    assert summary.flexure_results(index="V9t").equals(summary.flexure_results(index=2))
    assert summary.shear_results(index=("", "V9a")).equals(summary.shear_results(index=1))
    with pytest.raises(IndexError):
        summary.flexure_results(index=3)
    with pytest.raises(IndexError):
        summary.flexure_results(index=0)
    with pytest.raises(SummaryInputError) as raised:
        summary.flexure_results(index="V10")
    assert raised.value.code == "unknown_index"
    with pytest.raises(SummaryInputError):
        summary.flexure_results(index=("P1", "V9a"))
    assert summary.labels == ["V9a", "V9t"]
    assert [node.section.label for node in summary.nodes] == ["V9a", "V9t"]


# ============================================================================
# DESIGN IS NODE.DESIGN()
# ============================================================================


def test_design_is_node_design(sample_steel: SteelBar) -> None:
    """The cases where designing the flexure and then the shear on their own left a section short.

    ACI 318-19 20x50, cc 25, H25 / ADN 420, Mu 150 kN·m, Vu 250 kN: 2Ø25 and
    1eØ10/11 sank the bars, DCRb,bot 1.0005 on the summary's own check.
    ``Node.design()`` gives 2Ø25 + 1Ø20, 0.787 -- also under a second,
    hogging combination. EN 1992-1-1 30x80 under 30 kN·m: 3.047 cm² against
    A_s,min 3.054; ``Node.design()`` gives 2Ø10 + 2Ø10 = 3.142 cm². ACI 25x60,
    cc 30, Vz 150, My 120: 6.03 cm² and 0.9965 before, 2Ø20 and 0.962 now.
    """
    h25 = Concrete_ACI_318_19(name="H25", f_c=25 * MPa)
    for rows in (
        [{"Label": "V1", "Comb.": "C1", "Vz": 250, "My": 150}],
        [{"Label": "V1", "Comb.": "C1", "Vz": 250, "My": 150}, {"Label": "V1", "Comb.": "C2", "Vz": 100, "My": -60}],
    ):
        summary = BeamSummary(h25, sample_steel, beams([{"Label": "V1"}], units=GEOMETRY_UNITS), forces(rows))
        row = summary.design().iloc[1]
        assert (row["n1_bot"], row["db1_bot"], row["n2_bot"], row["db2_bot"]) == (2, 25, 1, 20)
        assert (row["legs"], row["dbs"], row["sl"]) == (2, 10, 11)
        check = summary.check().iloc[1]
        assert (check["DCRb,bot"], check[VERDICT_COLUMN]) == (0.787, PASS_MARK)

    en = Concrete_EN_1992_2004(name="C25/30", f_c=25 * MPa)
    b500 = SteelBar(name="B500S", f_y=500 * MPa)
    summary = BeamSummary(
        en,
        b500,
        beams([{"Label": "V1"}], units=GEOMETRY_UNITS, b=30, h=80),
        forces([{"Label": "V1", "Comb.": "C1", "My": 30}]),
    )
    row = summary.design().iloc[1]
    assert (row["n1_bot"], row["db1_bot"], row["n2_bot"], row["db2_bot"]) == (2, 10, 2, 10)
    assert summary.results[0].warnings == () and summary.results[0].passes

    summary = BeamSummary(
        h25,
        sample_steel,
        beams([{"Label": "V1"}], units=GEOMETRY_UNITS, b=25, h=60, cc=30),
        forces([{"Label": "V1", "Comb.": "C1", "Vz": 150, "My": 120}]),
    )
    row = summary.design().iloc[1]
    assert (row["n1_bot"], row["db1_bot"], row["n2_bot"]) == (2, 20, 0)
    assert summary.check().iloc[1]["DCRb,bot"] == 0.962


_CODES = {
    "ACI": (Concrete_ACI_318_19(name="C25", f_c=25 * MPa), SteelBar(name="ADN 420", f_y=420 * MPa)),
    "CIRSOC": (Concrete_CIRSOC_201_25(name="H30", f_c=30 * MPa), SteelBar(name="ADN 420", f_y=420 * MPa)),
    "EN": (Concrete_EN_1992_2004(name="C25/30", f_c=25 * MPa), SteelBar(name="B500S", f_y=500 * MPa)),
}

_CORPUS = [
    ({"b": 20, "h": 50}, [(60, 45, 0), (-110, -70, 10)]),
    ({"b": 25, "h": 60}, [(150, 120, 0)]),
    ({"b": 20, "h": 40}, [(80, -60, 0), (10, 170, 0)]),
    ({"b": 30, "h": 70}, [(200, 380, -50)]),
    ({"b": 20, "h": 50}, [(5, 30, 0)]),
]


@pytest.mark.parametrize("code", list(_CODES))
def test_design_matches_node_design_on_a_corpus(code: str) -> None:
    """Every section of a corpus is designed, warned and checked as ``Node.design()`` does by hand."""
    concrete, steel = _CODES[code]
    sections = beams([{"Label": f"V{i}", **geometry} for i, (geometry, _) in enumerate(_CORPUS)], units=GEOMETRY_UNITS)
    rows = [
        {"Label": f"V{i}", "Comb.": f"C{j}", "Vz": v, "My": m, "Nx": n}
        for i, (_, combos) in enumerate(_CORPUS)
        for j, (v, m, n) in enumerate(combos)
    ]
    summary = BeamSummary(concrete, steel, sections, forces(rows))
    summary.design()
    for i, (geometry, combos) in enumerate(_CORPUS):
        beam = RectangularBeam(
            label=f"V{i}",
            concrete=concrete,
            steel_bar=steel,
            width=geometry["b"] * cm,
            height=geometry["h"] * cm,
            c_c=25 * mm,
        )
        node = Node(
            beam, [Forces(label=f"C{j}", V_z=v * kN, M_y=m * kNm, N_x=n * kN) for j, (v, m, n) in enumerate(combos)]
        )
        node.design()
        ours = summary.nodes[i].section
        assert ours._bar_groups("bot") == beam._bar_groups("bot")
        assert ours._bar_groups("top") == beam._bar_groups("top")
        assert (ours._stirrup_n, ours._stirrup_d_b, ours._stirrup_s_l) == (
            beam._stirrup_n,
            beam._stirrup_d_b,
            beam._stirrup_s_l,
        )
        node.check()
        assert summary.results[i].warnings == node.warnings
        assert summary.results[i].bottom.DCR == max(c.bottom.DCR for c in beam.flexure_checks)


def test_design_is_a_function_of_its_inputs(h25: Any, sample_steel: SteelBar, tmp_path: Path) -> None:
    """Designing twice, after a round trip through the file, or from other bars, gives the same table."""
    summary = BeamSummary(h25, sample_steel, *geometry_only())
    first = summary.design()
    assert summary.design().equals(first)
    summary.to_excel(tmp_path / "beams.xlsx")
    summary.import_design(tmp_path / "beams.xlsx")
    assert summary.design().equals(first)

    other = beams(
        [
            {"Label": "V9a", "n1_top": 4, "db1_top": 32, "n1_bot": 2, "db1_bot": 25},
            {"Label": "V9t", "n1_top": 4, "db1_top": 32, "n1_bot": 2, "db1_bot": 25},
        ],
        b=20,
        h=40,
        legs=2,
        dbs=6,
        sl=30,
    )
    assert BeamSummary(h25, sample_steel, other, geometry_only()[1]).design().equals(first)


def test_a_sagging_only_beam_has_a_bare_top(sample_concrete: Any, sample_steel: SteelBar) -> None:
    """``Node.design()`` leaves no bars on top under positive moments only, and puts some below under negative ones.

    The summary is ``Node.design()``, so it inherits the asymmetry rather
    than add hanger bars of its own.
    """
    table = beams([{"Label": "P"}, {"Label": "N"}], units=GEOMETRY_UNITS)
    rows = [{"Label": "P", "Comb.": "C", "Vz": 30, "My": 60}, {"Label": "N", "Comb.": "C", "Vz": 30, "My": -60}]
    summary = BeamSummary(sample_concrete, sample_steel, table, forces(rows))
    summary.design()
    check = summary.check()
    assert list(check.iloc[1][["As,top", "As,bot", "Av"]]) == ["-", "2Ø16", "1sØ10/22"]
    assert list(check.iloc[2][["As,top", "As,bot"]]) == ["2Ø16", "2Ø12"]


# ============================================================================
# THE VERDICT COUNTS THE WARNINGS
# ============================================================================


def test_warnings_fail_the_verdict(sample_steel: SteelBar) -> None:
    """A section that misses a detailing limit is not OK, whatever its DCRs.

    EN 30x80 given 2Ø12 + 1Ø10 under 30 kN·m: below A_s,min. ACI 20x50 with
    stirrups 50 cm apart: past the spacing of §9.7.6.2.2.
    """
    en = Concrete_EN_1992_2004(name="C25/30", f_c=25 * MPa)
    summary = BeamSummary(
        en,
        SteelBar(name="B500S", f_y=500 * MPa),
        beams(
            [{"Label": "V1", "n1_bot": 2, "db1_bot": 12, "n2_bot": 1, "db2_bot": 10}], b=30, h=80, legs=2, dbs=6, sl=20
        ),
        forces([{"Label": "V1", "Comb.": "C1", "My": 30}]),
    )
    row = summary.check().iloc[1]
    assert row["DCRb,bot"] < 1 and row[VERDICT_COLUMN] == FAIL_MARK
    assert row["Warnings"] == "As_below_min (bottom)"

    aci = Concrete_ACI_318_19(name="H25", f_c=25 * MPa)
    summary = BeamSummary(
        aci,
        sample_steel,
        beams([{"Label": "V1", "n1_bot": 3, "db1_bot": 16}], legs=2, dbs=10, sl=50),
        forces([{"Label": "V1", "Comb.": "C1", "Vz": 60, "My": 50}]),
    )
    row = summary.check().iloc[1]
    assert row["DCRv"] < 1 and row[VERDICT_COLUMN] == FAIL_MARK
    assert "stirrup_spacing_exceeds_max (l)" in row["Warnings"]


def test_the_pass_threshold_is_shared() -> None:
    """A DCR of exactly 1, or 1 plus a rounding error, passes in the three summaries."""
    from mento.summary_base import verdict_passes

    for dcr in (1.0, 1.0 + 1e-12):
        assert verdict_passes((GoverningDemand(("C",), None, None, dcr),), ())
    assert not verdict_passes((GoverningDemand(("C",), None, None, 1.001),), ())
    assert not verdict_passes((GoverningDemand(("C",), None, None, 0.5, admissible=False),), ())


def test_unreinforced_sections_have_a_row(sample_concrete: Any, sample_steel: SteelBar) -> None:
    """No bars on either face: "no reinforcement: run design()", not a DCR of 10000.

    A section with some bars is checked, and a face in tension with none
    takes mento's sentinel and fails with ``As_below_min``.
    """
    sections = beams([{"Label": "bare"}, {"Label": "top-only", "n1_top": 2, "db1_top": 12}])
    rows = [
        {"Label": "bare", "Comb.": "C", "Vz": 10, "My": 20},
        {"Label": "top-only", "Comb.": "C", "Vz": 10, "My": 20},
    ]
    summary = BeamSummary(sample_concrete, sample_steel, sections, forces(rows))
    table = summary.check()
    assert table.iloc[1][VERDICT_COLUMN] == "no reinforcement: run design()"
    assert math.isnan(table.iloc[1]["DCRb,bot"])
    assert summary.results[0].status == "no_reinforcement"
    assert table.iloc[2][VERDICT_COLUMN] == FAIL_MARK
    assert "As_below_min (bottom)" in table.iloc[2]["Warnings"]


# ============================================================================
# READING THE TABLES
# ============================================================================


def test_labels_are_normalised(h25: Any, sample_steel: SteelBar, tmp_path: Path) -> None:
    """'V4 ' in the forces finds 'V4'; 4 and '4' are one label; a numeric label goes and comes back as text."""
    summary = BeamSummary(
        h25,
        sample_steel,
        beams([{"Label": "V4"}], units=GEOMETRY_UNITS),
        forces([{"Label": "V4 ", "Comb.": "C", "My": 1}]),
    )
    assert len(summary.nodes[0].forces) == 1
    with pytest.raises(SummaryInputError) as raised:
        BeamSummary(h25, sample_steel, beams([{"Label": 4}, {"Label": "4"}], units=GEOMETRY_UNITS), forces([]))
    assert raised.value.code == "duplicate_section"

    numeric = BeamSummary(
        h25,
        sample_steel,
        beams([{"Label": 101}], units=GEOMETRY_UNITS),
        forces([{"Label": 101.0, "Comb.": 1, "My": 5}]),
    )
    assert numeric.labels == ["101"]
    numeric.to_excel(tmp_path / "numeric.xlsx")
    back = BeamSummary.from_excel(h25, sample_steel, tmp_path / "numeric.xlsx")
    assert back.labels == ["101"]
    assert back.forces_table.iloc[1]["Comb."] == "1"


def test_a_level_is_part_of_the_key(h25: Any, sample_steel: SteelBar) -> None:
    """ETABS repeats B12 on every storey: (Level, Label) tells them apart."""
    units = {"Level": "", **GEOMETRY_UNITS}
    sections = beams([{"Level": "P1", "Label": "B12"}, {"Level": "P2", "Label": "B12"}], units=units)
    rows = [
        {"Level": "P1", "Label": "B12", "Comb.": "C", "Vz": 30, "My": 40},
        {"Level": "P2", "Label": "B12", "Comb.": "C", "Vz": 30, "My": -40},
    ]
    summary = BeamSummary(
        h25,
        sample_steel,
        sections,
        forces(rows, {"Level": "", "Label": "", "Comb.": "", "Nx": "kN", "Vz": "kN", "My": "kNm"}),
    )
    summary.design()
    assert summary.labels == [("P1", "B12"), ("P2", "B12")]
    check = summary.check()
    assert list(check.columns[:2]) == ["Level", "Beam"]
    assert list(check["Level"][1:]) == ["P1", "P2"]
    assert summary.flexure_results(index=("P2", "B12")).iloc[1]["Mu"] == -40.0
    with pytest.raises(SummaryInputError):
        summary.flexure_results(index="B12")  # two sections answer to it
    # Forces without a Level column name a label that is unique.
    one = BeamSummary(h25, sample_steel, sections.iloc[:2], forces([{"Label": "B12", "Comb.": "C", "My": 4}]))
    assert len(one.nodes[0].forces) == 1


def test_a_missing_axial_column_is_taken_as_zero_with_a_warning(h25: Any, sample_steel: SteelBar) -> None:
    rows = forces(
        [{"Label": "V1", "Comb.": "C", "Vz": 10, "My": 5}], {"Label": "", "Comb.": "", "Vz": "kN", "My": "kNm"}
    )
    with pytest.warns(SummaryInputWarning) as caught:
        summary = BeamSummary(h25, sample_steel, beams([{"Label": "V1"}], units=GEOMETRY_UNITS), rows)
    assert [w.message.code for w in caught] == ["no_axial_column"]
    assert summary.nodes[0].forces[0]._N_x.magnitude == 0


def test_an_axial_load_past_a_beams_is_flagged(h25: Any, sample_steel: SteelBar) -> None:
    """Pu >= 0.10 f'c Ag in compression (250 kN on a 20x50 in H25): ACI 318-19 / CIRSOC 201-25 §9.5.2.2.

    From there the moment strength is computed with the axial load (§22.4),
    which mento's beam does not do. A tension is not held to the limit by the
    clause, and EN 1992-1-1 states none for a beam.
    """
    rows = [
        {"Label": "V1", "Comb.": "C1", "Nx": 249, "My": 5},
        {"Label": "V1", "Comb.": "C2", "Nx": 250, "My": 5},
        {"Label": "V1", "Comb.": "C3", "Nx": -400, "My": 5},
    ]
    with pytest.warns(SummaryInputWarning) as caught:
        BeamSummary(h25, sample_steel, beams([{"Label": "V1"}], units=GEOMETRY_UNITS), forces(rows))
    beyond = [w.message for w in caught if w.message.code == "axial_load_beyond_beam"]
    assert len(beyond) == 1 and "'V1 / C2'" in str(beyond[0])
    assert "C1" not in str(beyond[0]) and "C3" not in str(beyond[0])
    assert "§9.5.2.2" in str(beyond[0])
    en = Concrete_EN_1992_2004(name="C25/30", f_c=25 * MPa)
    with warnings.catch_warnings(record=True) as caught_en:
        warnings.simplefilter("always")
        BeamSummary(en, sample_steel, beams([{"Label": "V1"}], units=GEOMETRY_UNITS), forces(rows))
    assert not [w for w in caught_en if getattr(w.message, "code", "") == "axial_load_beyond_beam"]


def test_forces_in_the_units_the_summaries_read(h25: Any, sample_steel: SteelBar) -> None:
    """kN and kNm, kip and kip·ft (or kipft): the units the single table read. Any other is named."""
    units = {"Label": "", "Comb.": "", "Nx": "kN", "Vz": "kN", "My": "kipft"}
    summary = BeamSummary(
        h25,
        sample_steel,
        beams([{"Label": "V1"}], units=GEOMETRY_UNITS),
        forces([{"Label": "V1", "Comb.": "C", "My": 3}], units),
    )
    assert summary.nodes[0].forces[0]._M_y.to("kip*ft").magnitude == pytest.approx(3)
    assert summary.forces_table.iloc[0]["My"] == "kipft"
    for unit in ("tf·m", "kN·m", "kNm/m"):
        with pytest.raises(SummaryInputError) as raised:
            BeamSummary(
                h25,
                sample_steel,
                beams([{"Label": "V1"}], units=GEOMETRY_UNITS),
                forces([{"Label": "V1", "Comb.": "C", "My": 3}], {**units, "My": unit}),
            )
        assert raised.value.code == "wrong_unit" and "'kNm', 'kip·ft', 'kipft'" in str(raised.value)


# ============================================================================
# FROM 1.4.0
# ============================================================================


def test_split_single_table_reproduces_1_4_0(
    sample_concrete: Any,
    sample_steel: SteelBar,
    sample_input_dataframe: pd.DataFrame,
    sample_input_with_nan: pd.DataFrame,
) -> None:
    """Each row a section of its own, the 2Ø8 1.4.0 placed written out: the DCRs of 1.4.0 come back.

    The DCRs below are those of 1.4.0 (85f9d35) on the same tables.
    """
    with pytest.warns(SummaryInputWarning) as caught:
        sections, rows = split_single_table(sample_input_dataframe, "beam")
    assert [w.message.code for w in caught] == ["second_layer_same_face"]
    summary = BeamSummary(sample_concrete, sample_steel, sections, rows)
    check = summary.check()
    dcrs = [tuple(check.iloc[i][["DCRb,top", "DCRb,bot", "DCRv"]]) for i in range(1, 5)]
    assert dcrs == [(0.0, 0.0, 0.59), (0.491, 0.0, 0.5), (0.0, 0.369, 1.047), (0.0, 1.165, 0.997)]
    # The face a row's moment does not pull gets the 2Ø8 1.4.0 placed, now written.
    assert list(sections.iloc[2][["n1_top", "db1_top", "n1_bot", "db1_bot"]]) == [2, 12, 2, 8.0]
    assert list(sections.iloc[1][["n1_top", "db1_top", "n3_bot", "db3_bot"]]) == [2, 8.0, 2, 12]
    # ns = 0 of 1.4.0 is legs = 0, which keeps the starter stirrup as 1.4.0 did: no stirrups, Av "-".
    assert (sections.iloc[1]["legs"], check.iloc[1]["Av"]) == (0, "-")
    assert sections.iloc[2]["legs"] == 2  # one closed stirrup, two legs

    with pytest.warns(SummaryInputWarning) as caught:
        sections, rows = split_single_table(sample_input_with_nan, "beam")
    codes = [w.message.code for w in caught]
    assert codes == ["second_layer_same_face", "dead_cells"]
    assert "V101: db2" in str(caught[1].message)
    check = BeamSummary(sample_concrete, sample_steel, sections, rows).check()
    assert [tuple(check.iloc[i][["DCRb,top", "DCRb,bot", "DCRv"]]) for i in (1, 2)] == [
        (0.0, 0.0, 0.624),
        (0.491, 0.0, 0.0),
    ]


def test_split_single_table_reproduces_the_1_4_0_snapshot() -> None:
    """V101 of the snapshot, ns = 0: DCRv 0.701 under ACI 318-19 and 0.539 under EN 1992-1-1, as in 1.4.0."""
    old = pd.DataFrame(
        {
            "Label": ["", "V101", "V102", "V103"],
            "Comb.": ["", "ELU 1", "ELU 2", "ELU 3"],
            "b": ["cm", 20, 20, 25],
            "h": ["cm", 50, 50, 60],
            "cc": ["mm", 25, 25, 25],
            "Nx": ["kN", 0, 0, 10],
            "Vz": ["kN", 20, -50, 100],
            "My": ["kNm", 0, -35, 60],
            "ns": ["", 0, 1.0, 1.0],
            "dbs": ["mm", 0, 6, 8],
            "sl": ["cm", 0, 20, 15],
            "n1": ["", 2.0, 2, 3.0],
            "db1": ["mm", 12, 12, 16],
            "n2": ["", 1.0, 1, 0.0],
            "db2": ["mm", 10, 16, 0],
            "n3": ["", 0.0, 0.0, 2.0],
            "db3": ["mm", 0, 0, 12],
            "n4": ["", 0, 0.0, 0],
            "db4": ["mm", 0, 0, 0],
        }
    )
    for concrete, steel, dcrv in (
        (Concrete_ACI_318_19(name="C25", f_c=25 * MPa), SteelBar(name="ADN 420", f_y=420 * MPa), [0.701, 0.5, 0.489]),
        (
            Concrete_EN_1992_2004(name="C25/30", f_c=25 * MPa),
            SteelBar(name="B500S", f_y=500 * MPa),
            [0.539, 0.391, 0.278],
        ),
    ):
        check = BeamSummary(concrete, steel, *split_single_table(old, "beam")).check()
        assert list(check["DCRv"][1:]) == dcrv


def test_split_single_table_handles_three_stations(sample_concrete: Any, sample_steel: SteelBar) -> None:
    """V1 at three stations, stirrups Ø8/10, Ø8/20, Ø8/10: three sections, V1, V1-2 and V1-3, with the DCRs of 1.4.0.

    The pull request's single table raised "its rows give different stirrups".
    """
    old = pd.DataFrame(
        {
            "Label": ["", "V1", "V1", "V1"],
            "Comb.": ["", "A", "B", "C"],
            "b": ["cm", 20, 20, 20],
            "h": ["cm", 50, 50, 50],
            "cc": ["mm", 25, 25, 25],
            "Nx": ["kN", 0, 0, 0],
            "Vz": ["kN", 100, 20, 100],
            "My": ["kNm", -80, 60, -50],
            "ns": ["", 1, 1, 1],
            "dbs": ["mm", 8, 8, 8],
            "sl": ["cm", 10, 20, 10],
            "n1": ["", 3, 2, 3],
            "db1": ["mm", 16, 12, 16],
            "n2": ["", 0, 0, 0],
            "db2": ["mm", 0, 0, 0],
            "n3": ["", 0, 0, 0],
            "db3": ["mm", 0, 0, 0],
            "n4": ["", 0, 0, 0],
            "db4": ["mm", 0, 0, 0],
        }
    )
    with pytest.warns(SummaryInputWarning) as caught:
        sections, rows = split_single_table(old, "beam")
    assert caught[0].message.code == "labels_renamed"
    assert "'V1' -> 'V1-2'" in str(caught[0].message)
    check = BeamSummary(sample_concrete, sample_steel, sections, rows).check()
    assert list(check["Beam"][1:]) == ["V1", "V1-2", "V1-3"]
    assert [tuple(check.iloc[i][["DCRb,top", "DCRb,bot", "DCRv"]]) for i in (1, 2, 3)] == [
        (0.818, 0.0, 0.49),
        (0.0, 1.56, 0.152),
        (0.511, 0.0, 0.49),
    ]


def test_a_second_layer_in_1_4_0_is_flagged(sample_input_dataframe: pd.DataFrame) -> None:
    """The 1.4.0 template drew n3/n4 as the top face, and 1.4.0 read them as a second layer of the tension face."""
    with pytest.warns(SummaryInputWarning) as caught:
        split_single_table(sample_input_dataframe, "beam")
    message = str(caught[0].message)
    assert caught[0].message.code == "second_layer_same_face"
    assert "'V101', 'V103'" in message


def test_split_single_table_names_rows_without_a_label(sample_concrete: Any, sample_steel: SteelBar) -> None:
    old = pd.DataFrame(
        {
            "Label": ["", None, "", "row-1"],
            "Comb.": ["", "A", "B", "C"],
            "b": ["cm", 20, 20, 20],
            "h": ["cm", 50, 50, 50],
            "cc": ["mm", 25, 25, 25],
            "Nx": ["kN", 0, 0, 0],
            "Vz": ["kN", 10, 10, 10],
            "My": ["kNm", 20, -20, 5],
            "ns": ["", 1, 0, 0],
            "dbs": ["mm", 8, 8, 0],
            "sl": ["cm", 20, 0, 0],
            "n1": ["", 2, 0, 2],
            "db1": ["mm", 12, 12, 10],
            "n2": ["", 0, 0, 0],
            "db2": ["mm", 0, 0, 0],
            "n3": ["", 0, 0, 0],
            "db3": ["mm", 0, 0, 0],
            "n4": ["", 0, 0, 0],
            "db4": ["mm", 0, 0, 0],
        }
    )
    with pytest.warns(SummaryInputWarning) as caught:
        sections, rows = split_single_table(old, "beam")
    assert [w.message.code for w in caught] == ["labels_renamed", "dead_cells"]
    # A row with no label takes its position, clear of a label already in the table.
    assert list(sections["Label"][1:]) == ["row-1-x", "row-2", "row-1"]
    # The second row's face had no bars: both faces get the 2Ø8 of 1.4.0.
    assert list(sections.iloc[2][["n1_top", "n1_bot", "legs", "dbs"]]) == [2, 2, 0, 0]
    assert len(BeamSummary(sample_concrete, sample_steel, sections, rows).nodes) == 3
    with pytest.raises(ValueError, match="'beam' or 'wall'"):
        split_single_table(old, "slab")  # type: ignore[arg-type]
    with pytest.raises(SummaryInputError) as raised:
        split_single_table(old.assign(extra=["", 1, 2, 3]), "beam")
    assert raised.value.code == "unknown_columns"


def test_split_single_table_in_us_customary_units() -> None:
    """The bar 1.4.0 placed is 2 #3 in inches; a bar column in mm keeps it in mm."""
    old = pd.DataFrame(
        {
            "Label": ["", "B1"],
            "Comb.": ["", "C"],
            "b": ["in", 12],
            "h": ["in", 24],
            "cc": ["in", 1.5],
            "Nx": ["kip", 0],
            "Vz": ["kip", 10],
            "My": ["kip·ft", 40],
            "ns": ["", 1],
            "dbs": ["in", 0.375],
            "sl": ["in", 8],
            "n1": ["", 3],
            "db1": ["mm", 19.05],
            "n2": ["", 0],
            "db2": ["mm", 0],
            "n3": ["", 0],
            "db3": ["mm", 0],
            "n4": ["", 0],
            "db4": ["mm", 0],
        }
    )
    sections, _ = split_single_table(old, "beam")
    assert sections.iloc[1]["db1_top"] == pytest.approx(9.525)
    assert sections.iloc[0]["db1_top"] == "mm"


# ============================================================================
# WRITING THE TABLES
# ============================================================================


@pytest.mark.parametrize("code", list(_CODES))
@pytest.mark.parametrize("shape", ["omitted", "zeros", "complete"])
def test_round_trip_is_exact(code: str, shape: str, tmp_path: Path) -> None:
    """The file gives back the tables, and after a design the same check and the same results.

    ``shape`` is the input: the reinforcement columns left out, given as
    zeros, or given in full.
    """
    concrete, steel = _CODES[code]
    if shape == "omitted":
        sections = beams([{"Label": "V1"}, {"Label": "V2"}], units=GEOMETRY_UNITS)
    elif shape == "zeros":
        sections = beams([{"Label": "V1"}, {"Label": "V2"}])
    else:
        sections = beams(
            [
                {"Label": "V1", "n1_top": 2, "db1_top": 12, "n1_bot": 3, "db1_bot": 16, "n3_bot": 2, "db3_bot": 12},
                {"Label": "V2", "n1_top": 2, "db1_top": 10, "n1_bot": 2, "db1_bot": 16, "n2_bot": 1, "db2_bot": 12},
            ],
            legs=2,
            dbs=8,
            sl=15,
        )
    rows = forces(
        [
            {"Label": "V1", "Comb.": "C1", "Vz": 90, "My": 110},
            {"Label": "V1", "Comb.": "C2", "Vz": -60, "My": -50, "Nx": 20},
            {"Label": "V2", "Comb.": "C1", "Vz": 40, "My": 60},
        ]
    )
    summary = BeamSummary(concrete, steel, sections, rows)
    path = tmp_path / "beams.xlsx"
    summary.to_excel(path)
    back = BeamSummary.from_excel(concrete, steel, path)
    for ours, theirs in zip(summary.tables(), back.tables()):
        pd.testing.assert_frame_equal(ours, theirs, check_dtype=False)

    summary.design()
    check = summary.check()
    summary.to_excel(path)
    back = BeamSummary.from_excel(concrete, steel, path)
    assert back.check().equals(check)
    assert back.results == summary.results
    for ours, theirs in zip(summary.nodes, back.nodes):
        assert ours.section._bar_groups("bot") == theirs.section._bar_groups("bot")
        assert ours.section._bar_groups("top") == theirs.section._bar_groups("top")
        assert ours.section._stirrup_s_l == theirs.section._stirrup_s_l


def test_legs_count_the_legs_of_closed_stirrups(sample_concrete: Any, sample_steel: SteelBar) -> None:
    """legs = 4 is two closed stirrups; an odd number is an error, as mento models closed stirrups only."""
    summary = BeamSummary(
        sample_concrete,
        sample_steel,
        beams([{"Label": "V1", "n1_bot": 3, "db1_bot": 16}], legs=4, dbs=8, sl=15),
        forces([]),
    )
    assert summary.nodes[0].section._stirrup_n == 2
    assert summary.sections_table.iloc[1]["legs"] == 4
    with pytest.raises(SummaryInputError) as raised:
        BeamSummary(sample_concrete, sample_steel, beams([{"Label": "V1"}], legs=3, dbs=8, sl=15), forces([]))
    assert raised.value.code == "odd_legs" and "'V1' is 3" in str(raised.value)
    with pytest.raises(SummaryInputError, match="'ns' -> legs: the number of stirrup legs") as raised:
        BeamSummary(sample_concrete, sample_steel, beams([{"Label": "V1", "ns": 1}], units=GEOMETRY_UNITS), forces([]))
    assert raised.value.code == "unknown_columns"


def test_section_row_is_the_inverse_of_section(sample_concrete: Any, sample_steel: SteelBar) -> None:
    """``_section(_section_row(s))`` rebuilds ``s``: the same groups and stirrups, a beam with none included."""
    summary = BeamSummary(
        sample_concrete,
        sample_steel,
        beams(
            [
                {"Label": "A", "legs": 4, "dbs": 10, "sl": 12, "n1_top": 2, "db1_top": 16, "n2_top": 1, "db2_top": 12},
                {"Label": "B", "n1_bot": 2, "db1_bot": 16, "n3_bot": 2, "db3_bot": 12, "n4_bot": 1, "db4_bot": 10},
            ]
        ),
        forces([]),
    )
    for node in summary.nodes:
        section = node.section
        again = summary._section(("", section.label), summary._section_row(section))
        for face in ("bot", "top"):
            assert again._bar_groups(face) == section._bar_groups(face)
        assert (again._stirrup_n, again._stirrup_d_b, again._stirrup_s_l) == (
            section._stirrup_n,
            section._stirrup_d_b,
            section._stirrup_s_l,
        )
    assert summary.nodes[1].section._stirrup_d_b == 8 * mm  # legs = 0 keeps the starter stirrup


def test_each_number_in_its_columns_unit(sample_concrete: Any, sample_steel: SteelBar, tmp_path: Path) -> None:
    """A design writes each value in the unit of its column: sl in mm, db in cm, b in m, moments in kip·ft."""
    units = {**GEOMETRY_UNITS, "b": "m", "legs": "", "dbs": "mm", "sl": "mm", "n1_bot": "", "db1_bot": "cm"}
    sections = beams([{"Label": "V1", "b": 0.2}], units=units)
    rows = forces(
        [{"Label": "V1", "Comb.": "C1", "Vz": 60, "My": 45}],
        {"Label": "", "Comb.": "", "Nx": "kN", "Vz": "kN", "My": "kip·ft"},
    )
    summary = BeamSummary(sample_concrete, sample_steel, sections, rows)
    designed = summary.design()
    assert list(designed.iloc[0][["b", "sl", "db1_bot", "db1_top"]]) == ["m", "mm", "cm", "mm"]
    row = designed.iloc[1]
    beam = summary.nodes[0].section
    assert row["b"] == 0.2
    assert row["sl"] == beam._stirrup_s_l.to("mm").magnitude
    assert row["db1_bot"] == beam._d_b1_b.to("cm").magnitude
    assert summary.forces_table.iloc[0]["My"] == "kip·ft"
    assert summary.forces_table.iloc[1]["My"] == 45
    summary.to_excel(tmp_path / "beams.xlsx")
    assert (
        BeamSummary.from_excel(sample_concrete, sample_steel, tmp_path / "beams.xlsx").check().equals(summary.check())
    )


def test_from_nodes_writes_what_the_nodes_are(sample_concrete: Any, sample_steel: SteelBar, tmp_path: Path) -> None:
    """Nodes built by hand, written as rows: the summary's check is each node's own, and so is the file's."""

    def beam(label: str) -> RectangularBeam:
        return RectangularBeam(
            label=label, concrete=sample_concrete, steel_bar=sample_steel, width=20 * cm, height=40 * cm, c_c=25 * mm
        )

    doubly = beam("D")
    doubly.set_transverse_rebar(1, 10 * mm, 17 * cm)
    doubly.set_longitudinal_rebar_top(3, 16 * mm)
    doubly.set_longitudinal_rebar_bot(2, 32 * mm)
    bare = beam("S")  # no stirrups
    bare.set_longitudinal_rebar_bot(3, 12 * mm)
    nodes = [
        Node(doubly, [Forces(label="C1", V_z=10 * kN, M_y=170 * kNm)]),
        Node(bare, [Forces(label="C1", V_z=15 * kN, M_y=30 * kNm)]),
    ]
    summary = BeamSummary.from_nodes(sample_concrete, sample_steel, nodes, units={"sl": "mm"})
    table = summary.check()
    for node in nodes:
        node.check()
    assert [r.warnings for r in summary.results] == [node.warnings for node in nodes]
    assert table.iloc[1]["DCRb,bot"] == round(max(c.bottom.DCR for c in doubly.flexure_checks), 3)
    sections = summary.sections_table
    assert sections.iloc[0]["sl"] == "mm" and sections.iloc[1]["sl"] == 170
    assert list(sections.iloc[2][["legs", "dbs", "sl"]]) == [0, 0, 0]

    summary.to_excel(tmp_path / "nodes.xlsx")
    assert BeamSummary.from_excel(sample_concrete, sample_steel, tmp_path / "nodes.xlsx").check().equals(table)


@pytest.mark.parametrize("gamma_s,epsilon_ud", [(1.0, None), (1.15, 0.01)])
def test_from_nodes_does_not_replace_a_steel_design_diagram(sample_concrete, sample_steel, gamma_s, epsilon_ud):
    other = SteelBar(name=sample_steel.name, f_y=sample_steel.f_y, gamma_s=gamma_s, epsilon_ud=epsilon_ud)
    beam = RectangularBeam(
        label="V1", concrete=sample_concrete, steel_bar=other, width=20 * cm, height=40 * cm, c_c=25 * mm
    )
    with pytest.raises(SummaryInputError, match="gamma_s") as raised:
        BeamSummary.from_nodes(sample_concrete, sample_steel, [Node(beam, [])])
    assert raised.value.code == "mixed_materials"


def test_from_nodes_does_not_replace_en_concrete_partial_parameters(sample_steel):
    source = Concrete_EN_1992_2004(name="C25", f_c=25 * MPa, alpha_cc=1.0)
    target = Concrete_EN_1992_2004(name="C25", f_c=25 * MPa, alpha_cc=0.85)
    beam = RectangularBeam(
        label="V1", concrete=source, steel_bar=sample_steel, width=20 * cm, height=40 * cm, c_c=25 * mm
    )
    with pytest.raises(SummaryInputError) as raised:
        BeamSummary.from_nodes(target, sample_steel, [Node(beam, [])])
    assert raised.value.code == "mixed_materials"


def test_from_nodes_rejects_what_a_row_cannot_hold(sample_concrete: Any, sample_steel: SteelBar) -> None:
    """Settings, an out-of-plane moment, other materials, a stirrup diameter without stirrups: not a row."""
    from mento import BeamSettings

    def beam(**kwargs: Any) -> RectangularBeam:
        options = dict(concrete=sample_concrete, steel_bar=sample_steel, width=20 * cm, height=40 * cm, c_c=25 * mm)
        options.update(kwargs)
        return RectangularBeam(label="V1", **options)

    cases = [
        (Node(beam(settings=BeamSettings(layers_spacing=40 * mm)), [Forces(M_y=10 * kNm)]), "node_not_representable"),
        (Node(beam(), [Forces(M_y=10 * kNm, M_x=5 * kNm)]), "node_not_representable"),
        (Node(beam(steel_bar=SteelBar(name="B500S", f_y=500 * MPa)), [Forces(M_y=10 * kNm)]), "mixed_materials"),
        (Node(beam(concrete=Concrete_ACI_318_19(name="H30", f_c=30 * MPa)), [Forces(M_y=10 * kNm)]), "mixed_materials"),
    ]
    stirrup = beam()
    stirrup._stirrup_d_b = 10 * mm  # a diameter set on a beam that carries no stirrups
    cases.append((Node(stirrup, [Forces(M_y=10 * kNm)]), "node_not_representable"))
    for node, code in cases:
        with pytest.raises(SummaryInputError) as raised:
            BeamSummary.from_nodes(sample_concrete, sample_steel, [node])
        assert raised.value.code == code
    assert "a row with legs = 0 does not hold" in str(raised.value)
    with pytest.raises(SummaryInputError) as raised:
        BeamSummary.from_nodes(sample_concrete, sample_steel, [Node(beam(), []), Node(beam(), [])])
    assert raised.value.code == "duplicate_section"
    with pytest.raises(SummaryInputError) as raised:
        BeamSummary.from_nodes(sample_concrete, sample_steel, [Node(beam(), [])], units={"sl": "kN"})
    assert raised.value.code == "wrong_unit"


def test_to_excel_accepts_a_buffer(h25: Any, sample_steel: SteelBar) -> None:
    """For the web: a buffer in, the same summary out."""
    summary = BeamSummary(h25, sample_steel, *one_continuous_section())
    buffer = io.BytesIO()
    summary.to_excel(buffer)
    buffer.seek(0)
    sheets = pd.read_excel(buffer, sheet_name=None)
    assert list(sheets) == ["Sections", "Forces"]
    buffer.seek(0)
    assert BeamSummary.from_excel(h25, sample_steel, buffer).check().equals(summary.check())


def test_to_excel_ignores_the_language(h25: Any, sample_steel: SteelBar, tmp_path: Path) -> None:
    """Sheets and headers stay in English in a Spanish session, and read back."""
    summary = BeamSummary(h25, sample_steel, *one_continuous_section())
    set_language("es")
    try:
        assert "Viga" in summary.check().columns
        summary.to_excel(tmp_path / "es.xlsx")
    finally:
        set_language("en")
    sheets = pd.read_excel(tmp_path / "es.xlsx", sheet_name=None)
    assert list(sheets) == ["Sections", "Forces"] and "Label" in sheets["Sections"].columns
    assert BeamSummary.from_excel(h25, sample_steel, tmp_path / "es.xlsx").check().equals(summary.check())


def test_notes_round_trip(h25: Any, sample_steel: SteelBar, tmp_path: Path) -> None:
    """Free text in Notes, in both sheets: not read into the calculation, written back as given."""
    sections, rows = one_continuous_section()
    sections = sections.assign(Notes=["", "pórtico 3, eje B"])
    rows = rows.assign(Notes=["", "apoyo izq.", ""])
    summary = BeamSummary(h25, sample_steel, sections, rows)
    summary.design()
    assert summary.sections_table.iloc[1]["Notes"] == "pórtico 3, eje B"
    summary.to_excel(tmp_path / "notes.xlsx")
    back = BeamSummary.from_excel(h25, sample_steel, tmp_path / "notes.xlsx")
    assert back.sections_table.iloc[1]["Notes"] == "pórtico 3, eje B"
    assert list(back.forces_table["Notes"][1:]) == ["apoyo izq.", ""]


def test_blank_rows_and_extra_sheets(h25: Any, sample_steel: SteelBar, tmp_path: Path) -> None:
    """A separator row left blank is skipped wherever it is, and sheets other than the two are ignored."""
    sections, rows = support_and_midspan()
    blank = pd.DataFrame([{column: None for column in sections.columns}])
    spaced = pd.concat([sections.iloc[:2], blank, sections.iloc[2:]], ignore_index=True)
    path = tmp_path / "book.xlsx"
    with pd.ExcelWriter(path) as writer:
        pd.DataFrame({"a": [1]}).to_excel(writer, sheet_name="Cover", index=False)
        spaced.to_excel(writer, sheet_name="Sections", index=False)
        rows.to_excel(writer, sheet_name="Forces", index=False)
    summary = BeamSummary.from_excel(h25, sample_steel, path)
    assert summary.labels == ["V9a", "V9t"]


def test_from_excel_with_other_sheet_names(h25: Any, sample_steel: SteelBar, tmp_path: Path) -> None:
    sections, rows = support_and_midspan()
    path = tmp_path / "book.xlsx"
    with pd.ExcelWriter(path) as writer:
        sections.to_excel(writer, sheet_name="Vigas", index=False)
        rows.to_excel(writer, sheet_name="Solicitaciones", index=False)
    summary = BeamSummary.from_excel(h25, sample_steel, path, sections_sheet="Vigas", forces_sheet="Solicitaciones")
    assert summary.labels == ["V9a", "V9t"]
    with pytest.raises(SummaryInputError) as raised:
        BeamSummary.from_excel(h25, sample_steel, path)
    assert raised.value.code == "missing_sheet"


# ============================================================================
# THE WORD REPORT
# ============================================================================


def _built_document(summary: Any, monkeypatch: pytest.MonkeyPatch, index: Any = 1) -> Any:
    """Build the Word summary and hand back the document instead of saving it."""
    captured: list = []
    monkeypatch.setattr(DocumentBuilder, "save", lambda self, *_: captured.append(self.doc))
    summary.results_detailed_doc(index=index)
    return captured[0]


def _rows(table: Any) -> list:
    return [[cell.text for cell in row.cells] for row in table.rows]


def _table_after(doc: Any, heading: str) -> Any:
    """The first table after the heading whose text is ``heading``."""
    from docx.table import Table

    found = False
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p") and "".join(t.text or "" for t in child.iter(qn("w:t"))) == heading:
            found = True
        elif found and child.tag == qn("w:tbl"):
            return Table(child, doc)
    raise AssertionError(f"no table after {heading!r}")


def _cell_fill(cell: Any) -> Optional[str]:
    tc_pr = cell._element.find(qn("w:tcPr"))
    if tc_pr is None:
        return None
    shd = tc_pr.find(qn("w:shd"))
    return None if shd is None else shd.get(qn("w:fill"))


def _table_width_cm(table: Any) -> float:
    return sum(Emu(cell.width).cm for cell in table.rows[0].cells if cell.width is not None)


def _usable_width_cm(doc: Any) -> float:
    section = doc.sections[0]
    return Emu(section.page_width - section.left_margin - section.right_margin).cm


def test_the_report_prints_check_rather_than_a_subset_of_it(
    beam_summary: BeamSummary, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The notebook and the Word report show the same summary, warnings and verdict included."""
    shown = beam_summary.check()
    table = _table_after(_built_document(beam_summary, monkeypatch), "Design Check Summary")
    assert _rows(table)[0] == list(shown.columns)
    assert _rows(table)[0][-2:] == ["Warnings", VERDICT_COLUMN]


def test_the_verdict_column_is_shaded(
    sample_concrete: Any, sample_steel: SteelBar, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Green for a pass, red for a fail; the header, the unit row and a section not checked, neither."""
    sections = beams(
        [
            {"Label": "passes", "b": 25, "n1_bot": 3, "db1_bot": 16, "legs": 2, "dbs": 8, "sl": 20},
            {"Label": "fails", "h": 30, "n1_bot": 2, "db1_bot": 8, "legs": 2, "dbs": 6, "sl": 12},
            {"Label": "idle", "n1_bot": 2, "db1_bot": 8},
        ]
    )
    rows = [
        {"Label": "passes", "Comb.": "C", "Vz": 40, "My": 40},
        {"Label": "fails", "Comb.": "C", "Vz": 250, "My": 200},
    ]
    summary = BeamSummary(sample_concrete, sample_steel, sections, forces(rows))
    table = _table_after(_built_document(summary, monkeypatch), "Design Check Summary")
    status = len(table.columns) - 1
    by_label = {row.cells[0].text: row.cells[status] for row in table.rows}
    assert _cell_fill(by_label["Beam"]) is None
    assert (by_label["passes"].text, _cell_fill(by_label["passes"])) == (PASS_MARK, "C6EFCE")
    assert (by_label["fails"].text, _cell_fill(by_label["fails"])) == (FAIL_MARK, "FFC7CE")
    assert (by_label["idle"].text, _cell_fill(by_label["idle"])) == ("no forces", None)


def test_the_all_beam_tables_are_set_a_point_smaller(
    beam_summary: BeamSummary, monkeypatch: pytest.MonkeyPatch
) -> None:
    doc = _built_document(beam_summary, monkeypatch)
    for heading in ("Flexure Results", "Shear Results", "Design Check Summary"):
        table = _table_after(doc, heading)
        sizes = {
            run.font.size.pt
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
            for run in paragraph.runs
            if run.font.size is not None
        }
        assert sizes == {float(SUMMARY_FONT_SIZE)}, heading


def _report_summaries(concrete: Any, steel: SteelBar, beam_summary: BeamSummary) -> dict:
    """A beam, a slab and a wall summary, each with a warning, a section without forces and a long label."""
    slab_sections = slabs(
        [
            {"Label": "L1-long-label", "db1_bot": 12, "s1_bot": 40},
            {"Label": "L2", "db1_bot": 10, "s1_bot": 20},
        ]
    )
    slab_forces = forces([{"Label": "L1-long-label", "Comb.": "1.2D+1.6L", "Vz": 20, "My": 10}])
    wall_sections = walls(
        [{"Level": "Level 1", "Label": "M1"}, {"Level": "Level 2", "Label": "M1"}], dbh=20, sh=50, dbv=10, sv=30
    )
    wall_rows = wall_forces([{"Level": "Level 1", "Label": "M1", "Comb.": "ELU 1", "Vz": 264, "My": -172}])
    return {
        "beam": beam_summary,
        "slab": OneWaySlabSummary(concrete, steel, slab_sections, slab_forces),
        "wall": ShearWallSummary(concrete, steel, wall_sections, wall_rows),
    }


@pytest.mark.parametrize("element", ["beam", "slab", "wall"])
@pytest.mark.parametrize("language", ["en", "es"])
def test_no_table_in_the_report_runs_past_the_page(
    sample_concrete: Any,
    sample_steel: SteelBar,
    beam_summary: BeamSummary,
    monkeypatch: pytest.MonkeyPatch,
    language: str,
    element: str,
) -> None:
    summary = _report_summaries(sample_concrete, sample_steel, beam_summary)[element]
    set_language(language)
    try:
        doc = _built_document(summary, monkeypatch)
    finally:
        set_language("en")
    usable = _usable_width_cm(doc)
    for i, table in enumerate(doc.tables):
        assert _table_width_cm(table) <= usable + 0.01, f"table {i} is wider than the text column"
    cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    assert "nan" not in cells


def test_word_sections_show_the_designed_bars(
    h25: Any, sample_steel: SteelBar, monkeypatch: pytest.MonkeyPatch
) -> None:
    """After design(), the sections table of the report lists what was designed, both faces, as check() writes it."""
    summary = BeamSummary(h25, sample_steel, *geometry_only())
    summary.design()
    rows = _rows(_table_after(_built_document(summary, monkeypatch), "Beam Sections"))
    assert rows[0] == ["Label", "b×h", "cc", "As,top", "As,bot", "Av"]
    assert rows[1] == ["", "cm", "mm", "", "", ""]
    check = summary.check()
    for row, (_, expected) in zip(rows[2:], check.iloc[1:].iterrows()):
        assert row == [expected["Beam"], "20×40", "25", expected["As,top"], expected["As,bot"], expected["Av"]]
    assert rows[3][3:5] == ["3Ø16", "2Ø32"]


def test_the_report_lists_the_forces_with_their_signs(
    h25: Any, sample_steel: SteelBar, monkeypatch: pytest.MonkeyPatch
) -> None:
    summary = BeamSummary(h25, sample_steel, *support_and_midspan())
    doc = _built_document(summary, monkeypatch)
    rows = _rows(_table_after(doc, "Forces"))
    assert rows[0] == ["Label", "Comb.", "Nx", "Vz", "My"]
    assert rows[1] == ["", "", "kN", "kN", "kNm"]
    assert rows[2:] == [["V9a", "apoyo", "0", "80", "-60"], ["V9t", "tramo", "0", "10", "170"]]
    assert any(p.text.startswith("Nx > 0 is compression") for p in doc.paragraphs)


@pytest.mark.parametrize("language", ["en", "es"])
def test_the_report_words_every_warning(
    h25: Any, sample_steel: SteelBar, monkeypatch: pytest.MonkeyPatch, language: str
) -> None:
    """Case A plus a section with no forces: each warning in full, with its face and its combinations."""
    sections, rows = support_and_midspan()
    sections = beams(
        [
            {"Label": "V9a", "n1_top": 3, "db1_top": 16, "n1_bot": 2, "db1_bot": 8},
            {"Label": "V9t", "n1_top": 2, "db1_top": 8, "n1_bot": 2, "db1_bot": 32},
            {"Label": "V10", "n1_bot": 2, "db1_bot": 12},
        ],
        b=20,
        h=40,
        legs=2,
        dbs=10,
        sl=17,
    )
    summary = BeamSummary(h25, sample_steel, sections, rows)
    set_language(language)
    try:
        doc = _built_document(summary, monkeypatch)
        heading = translate("Warnings")
        expected_messages = [w.message for w in summary.warnings[("", "V9t")]]
        no_forces = translate("no forces")
        face = translate("Bottom")
    finally:
        set_language("en")
    table = _rows(_table_after(doc, heading))
    assert [row[1:3] for row in table[1:]] == [[face, "tramo"], ["-", "-"], ["-", "-"]]
    assert [row[0] for row in table[1:]] == ["V9t", "V9t", "V10"]
    assert [row[3] for row in table[1:]] == [*expected_messages, no_forces]


def test_a_report_without_warnings_says_so(h25: Any, sample_steel: SteelBar, monkeypatch: pytest.MonkeyPatch) -> None:
    """Case B passes with no warning: the Warnings heading is followed by a sentence, not an empty table."""
    doc = _built_document(BeamSummary(h25, sample_steel, *one_continuous_section()), monkeypatch)
    texts = [p.text for p in doc.paragraphs]
    assert texts[texts.index("Warnings") + 1] == "No section misses a detailing limit."


def test_table_widths_are_honoured_and_never_padded(beam_summary: BeamSummary, monkeypatch: pytest.MonkeyPatch) -> None:
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        doc = _built_document(beam_summary, monkeypatch)
    assert not [str(w.message) for w in caught if "widths were given" in str(w.message)]
    for table in doc.tables:
        assert table.autofit is False
        layout = table._tbl.find(qn("w:tblPr")).find(qn("w:tblLayout"))
        assert layout is not None and layout.get(qn("w:type")) == "fixed"


def test_the_shear_summary_drops_the_capacity_ticks(beam_summary: BeamSummary, monkeypatch: pytest.MonkeyPatch) -> None:
    header = _rows(_table_after(_built_document(beam_summary, monkeypatch), "Shear Results"))[0]
    assert "Vu≤ØVn" not in header and "Vu≤ØVmax" not in header
    assert {"ØVn", "ØVmax", "DCR"} <= set(header)
    assert {"Vu≤ØVn", "Vu≤ØVmax"} <= set(beam_summary.shear_results().columns)


@pytest.mark.parametrize(
    "concrete, tick, demand, capacity",
    [
        (Concrete_ACI_318_19(name="H25", f_c=25 * MPa), "Mu≤ØMn", "Mu", "ØMn"),
        (Concrete_EN_1992_2004(name="C25/30", f_c=25 * MPa), "MEd≤MRd", "MEd", "MRd"),
    ],
    ids=["ACI", "EN"],
)
def test_the_flexure_summary_drops_the_codes_own_capacity_tick(
    concrete: Any,
    tick: str,
    demand: str,
    capacity: str,
    sample_steel: SteelBar,
    sample_input_dataframe: pd.DataFrame,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    summary = BeamSummary(concrete, sample_steel, *split_single_table(sample_input_dataframe, "beam"))
    header = _rows(_table_after(_built_document(summary, monkeypatch), "Flexure Results"))[0]
    assert tick not in header and {demand, capacity, "DCR"} <= set(header)
    assert tick in summary.flexure_results().columns
    check = _rows(_table_after(_built_document(summary, monkeypatch), "Design Check Summary"))[0]
    assert {f"{demand},top", f"{demand},bot"} <= set(check)


def test_limit_check_verdicts_are_shaded(
    sample_concrete: Any, sample_steel: SteelBar, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Ø10 stirrups at 40 cm on a 25x50 are past d/2: that row is red, the minimum shear reinforcement green."""
    summary = BeamSummary(
        sample_concrete,
        sample_steel,
        beams([{"Label": "sparse", "b": 25, "n1_bot": 3, "db1_bot": 16}], legs=2, dbs=10, sl=40),
        forces([{"Label": "sparse", "Comb.": "C", "Vz": 72, "My": 90}]),
    )
    doc = _built_document(summary, monkeypatch)
    limit_tables = [t for t in doc.tables if t.rows[0].cells[0].text == "Check"]
    by_check = {r.cells[0].text: r.cells[len(t.columns) - 1] for t in limit_tables for r in t.rows[1:]}
    assert (by_check["Stirrup spacing along length"].text, _cell_fill(by_check["Stirrup spacing along length"])) == (
        FAIL_MARK,
        "FFC7CE",
    )
    assert _cell_fill(by_check["Minimum shear reinforcement"]) == "C6EFCE"


def test_the_report_of_a_section_by_label(beam_summary: BeamSummary, monkeypatch: pytest.MonkeyPatch) -> None:
    doc = _built_document(beam_summary, monkeypatch, index="V103")
    assert any("V103" in p.text for p in doc.paragraphs)
    with pytest.raises(IndexError):
        beam_summary.results_detailed_doc(index=0)


def test_a_section_that_only_just_passes_is_not_reported_as_failing(
    sample_concrete: Any, sample_steel: SteelBar
) -> None:
    """Shear DCR 0.997 shows as 1.00 at two decimals; the verdict reads the DCR."""
    summary = BeamSummary(
        sample_concrete,
        sample_steel,
        beams([{"Label": "edge", "b": 25, "n1_bot": 3, "db1_bot": 16}], legs=2, dbs=8, sl=20),
        forces([{"Label": "edge", "Comb.": "C", "Vz": 145.392, "My": 10}]),
    )
    row = summary.check().iloc[1]
    assert 0.995 <= summary.results[0].shear.DCR < 1.0
    assert row[VERDICT_COLUMN] == PASS_MARK


def test_a_beam_that_is_not_tension_controlled_fails_the_summary(sample_concrete: Any, sample_steel: SteelBar) -> None:
    """25x40, 3Ø25 ++ 3Ø25 below under 100 kN·m: every DCR below 1, and still ❌ (§9.3.3.1)."""
    summary = BeamSummary(
        sample_concrete,
        sample_steel,
        beams(
            [{"Label": "over", "b": 25, "h": 40, "n1_bot": 3, "db1_bot": 25, "n3_bot": 3, "db3_bot": 25}],
            legs=2,
            dbs=10,
            sl=15,
        ),
        forces([{"Label": "over", "Comb.": "C", "Vz": 20, "My": 100}]),
    )
    row = summary.check().iloc[1]
    assert max(row["DCRb,bot"], row["DCRv"]) < 1.0
    assert summary.results[0].bottom.admissible is False
    assert row[VERDICT_COLUMN] == FAIL_MARK


def test_the_check_is_translated(h25: Any, sample_steel: SteelBar) -> None:
    summary = BeamSummary(h25, sample_steel, *support_and_midspan())
    set_language("es")
    try:
        table = summary.check()
    finally:
        set_language("en")
    assert {"Viga", "Advertencias", "¿Ok?"} <= set(table.columns)
    # The codes are mento's, the same in every language.
    assert (
        table.iloc[2]["Advertencias"] == "not_tension_controlled (bottom), stirrup_spacing_exceeds_compression_support"
    )


def test_print_follows_the_language(h25: Any, sample_steel: SteelBar, tmp_path: Path, capsys: Any) -> None:
    summary = BeamSummary(h25, sample_steel, *geometry_only())
    set_language("es")
    try:
        summary.design()
        summary.export_design(tmp_path / "x.xlsx")
        summary.import_design(tmp_path / "x.xlsx")
    finally:
        set_language("en")
    out = capsys.readouterr().out
    assert "Diseño completo" in out and "escritas en" in out and "leídas de" in out


def test_deprecated_summary_module_still_exports_beam_summary() -> None:
    """mento.summary is the pre-rename import path and must keep working."""
    import importlib
    import sys

    sys.modules.pop("mento.summary", None)
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        legacy = importlib.import_module("mento.summary")
    assert legacy.BeamSummary is BeamSummary
    assert any(issubclass(w.category, DeprecationWarning) for w in caught)


def test_a_copy_of_the_summary_is_independent(h25: Any, sample_steel: SteelBar) -> None:
    summary = BeamSummary(h25, sample_steel, *geometry_only())
    clone = copy.deepcopy(summary)
    clone.design()
    assert summary.sections_table.iloc[1]["n1_top"] == 0
    assert clone.sections_table.iloc[1]["n1_top"] == 2
